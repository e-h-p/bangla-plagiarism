# -*- coding: utf-8 -*-
"""Step 4 - BanglaBERT evaluation on the raw and polished BanglaPlagPair-1K.

BanglaBERT evaluation on the updated
dataset. Two encoders, two dataset states:

    csebuetnlp/banglabert           -> BanglaBERT_Cosine_BUETCSE
    sagorsarker/bangla-bert-base    -> BanglaBERT_Cosine_SagorSarkar

Methodology (unchanged from the reference):
  * mean pooling over the last hidden state, attention-mask weighted
  * L2 normalization of the pooled vector
  * pair score = dot product of the two normalized vectors (= cosine similarity)
  * max_length 512, batch size 8, CPU, model in eval() mode
  * unique texts are encoded once and reused (identical text -> identical vector)
  * fixed threshold 0.50 plus the 0.10-0.90 sweep
  * best threshold chosen by F1 -> Accuracy -> Recall -> Precision

Output structure:
    banglabert_evaluation/
      normal/    <- raw dataset
      polished/  <- polished dataset
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from transformers import AutoModel, AutoTokenizer

from pipeline_common import (
    BANGLABERT_DIR,
    METRICS,
    POLISHED_SCORES_XLSX,
    RAW_SCORES_XLSX,
    SEED,
    VALIDATION_DIR,
    banner,
    best_thresholds,
    classification_metrics,
    ensure_dirs,
    set_global_seed,
    threshold_sweep,
    THRESHOLDS,
)


MODELS = [
    ("csebuetnlp/banglabert", "BanglaBERT_Cosine_BUETCSE"),
    ("sagorsarker/bangla-bert-base", "BanglaBERT_Cosine_SagorSarkar"),
]

TRADITIONAL_METRICS = list(METRICS)  # Cosine, Jaccard, LCS, Sequence, EditDist, WordOverlap

ID_COLUMNS = [
    "pair_no",
    "pair_id",
    "variant_type",
    "pair_type",
    "binary_label",
    "source_link_text1",
    "source_link_text2",
]


# ═══════════════════════════════════════════════════════════════════
# 1. Embeddings
# ═══════════════════════════════════════════════════════════════════

def mean_pool(last_hidden_state: torch.Tensor, attention_mask: torch.Tensor) -> torch.Tensor:
    """Mean pooling over token embeddings, respecting the attention mask."""
    mask = attention_mask.unsqueeze(-1).expand(last_hidden_state.size()).float()
    summed = torch.sum(last_hidden_state * mask, dim=1)
    counts = torch.clamp(mask.sum(dim=1), min=1e-9)
    return summed / counts


def encode_unique_texts(
    texts: list[str],
    *,
    model_name: str,
    batch_size: int,
    max_length: int,
    device: str,
) -> tuple[dict[str, np.ndarray], dict[str, object]]:
    """Encode unique texts into L2-normalized embeddings."""
    print(f"  Loading model: {model_name} ...")
    tokenizer = AutoTokenizer.from_pretrained(model_name)
    model = AutoModel.from_pretrained(model_name)
    model.eval()
    model.to(device)

    unique_texts = list(dict.fromkeys(texts))
    embeddings: dict[str, np.ndarray] = {}
    total = len(unique_texts)
    truncated = 0
    token_lengths: list[int] = []

    print(f"  Encoding {total} unique texts (batch_size={batch_size}, max_length={max_length}) ...")
    with torch.no_grad():
        for start in range(0, total, batch_size):
            batch_texts = unique_texts[start : start + batch_size]

            # Diagnostic: how many texts hit the 512-token ceiling?
            raw_lengths = [len(tokenizer(t, truncation=False)["input_ids"]) for t in batch_texts]
            token_lengths.extend(raw_lengths)
            truncated += sum(1 for length in raw_lengths if length > max_length)

            encoded = tokenizer(
                batch_texts,
                padding=True,
                truncation=True,
                max_length=max_length,
                return_tensors="pt",
            )
            encoded = {key: value.to(device) for key, value in encoded.items()}
            outputs = model(**encoded)
            pooled = mean_pool(outputs.last_hidden_state, encoded["attention_mask"])
            normalized = torch.nn.functional.normalize(pooled, p=2, dim=1)
            for text, vector in zip(batch_texts, normalized.cpu().numpy(), strict=True):
                embeddings[text] = vector.astype(np.float32)

            done = min(start + batch_size, total)
            print(f"    [{done}/{total}] texts encoded", end="\r")

    print()
    stats = {
        "unique_texts": total,
        "truncated_at_max_length": truncated,
        "token_length_min": int(min(token_lengths)) if token_lengths else 0,
        "token_length_mean": round(float(np.mean(token_lengths)), 2) if token_lengths else 0.0,
        "token_length_max": int(max(token_lengths)) if token_lengths else 0,
        "embedding_dim": int(next(iter(embeddings.values())).shape[0]) if embeddings else 0,
    }
    return embeddings, stats


def compute_banglabert_scores(
    df: pd.DataFrame,
    *,
    model_name: str,
    cosine_column: str,
    batch_size: int,
    max_length: int,
    device: str,
) -> tuple[pd.DataFrame, np.ndarray, np.ndarray, dict[str, object]]:
    texts = [str(v) for v in df["text1"].tolist()] + [str(v) for v in df["text2"].tolist()]
    embeddings, stats = encode_unique_texts(
        texts,
        model_name=model_name,
        batch_size=batch_size,
        max_length=max_length,
        device=device,
    )

    scored = df.copy()
    text1_vectors = [embeddings[str(t)] for t in scored["text1"].tolist()]
    text2_vectors = [embeddings[str(t)] for t in scored["text2"].tolist()]
    cosine_scores = [float(np.dot(v1, v2)) for v1, v2 in zip(text1_vectors, text2_vectors, strict=True)]

    scored[cosine_column] = [round(s, 6) for s in cosine_scores]

    t1_arr = np.stack(text1_vectors) if text1_vectors else np.empty((0, 0), dtype=np.float32)
    t2_arr = np.stack(text2_vectors) if text2_vectors else np.empty((0, 0), dtype=np.float32)

    stats = dict(stats)
    stats.update(
        {
            "model": model_name,
            "score_column": cosine_column,
            "cosine_min": round(float(np.min(cosine_scores)), 6),
            "cosine_mean": round(float(np.mean(cosine_scores)), 6),
            "cosine_max": round(float(np.max(cosine_scores)), 6),
            "cosine_std": round(float(np.std(cosine_scores, ddof=1)), 6),
            "l2_norm_max_deviation": round(
                float(max(abs(np.linalg.norm(t1_arr, axis=1) - 1).max(),
                          abs(np.linalg.norm(t2_arr, axis=1) - 1).max())),
                8,
            ),
        }
    )
    return scored, t1_arr, t2_arr, stats


# ═══════════════════════════════════════════════════════════════════
# 2. Interactive workbook (verbatim from the reference script)
# ═══════════════════════════════════════════════════════════════════

def excel_safe(value: object) -> object:
    if pd.isna(value):
        return ""
    return value


def build_interactive_workbook(
    labeled_scores: pd.DataFrame,
    metrics: list[str],
    output_file: Path,
    source_description: str = "",
) -> None:
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "dashboard"
    scores_sheet = wb.create_sheet("scores")

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    buetcse_fill = PatternFill(fill_type="solid", fgColor="FCE5CD")
    sagorsarkar_fill = PatternFill(fill_type="solid", fgColor="D5F5E3")
    traditional_fill = PatternFill(fill_type="solid", fgColor="E8DAEF")
    bold = Font(bold=True)

    dashboard["A1"] = "Threshold"
    dashboard["A1"].font = bold
    dashboard["B1"] = 0.5
    dashboard["B1"].number_format = "0.00"
    dashboard["K1"] = "Change B1 manually to observe result changes"
    dashboard["K1"].font = Font(italic=True)
    if source_description:
        dashboard["K2"] = source_description
        dashboard["K2"].font = Font(italic=True)

    dashboard["K4"] = "Legend:"
    dashboard["K4"].font = bold
    dashboard["L4"] = "BUETCSE BanglaBERT"
    dashboard["L4"].fill = buetcse_fill
    dashboard["M4"] = "SagorSarkar BanglaBERT"
    dashboard["M4"].fill = sagorsarkar_fill
    dashboard["N4"] = "Traditional Metrics"
    dashboard["N4"].fill = traditional_fill

    metrics_header_row = 4
    for col_idx, header in enumerate(
        ["Metric", "TP", "FP", "TN", "FN", "Accuracy", "Precision", "Recall", "F1"], start=1
    ):
        cell = dashboard.cell(metrics_header_row, col_idx, header)
        cell.font = bold
        cell.fill = header_fill

    score_headers = labeled_scores.columns.tolist()
    metric_columns: dict[str, str] = {}

    for col_idx, header in enumerate(score_headers, start=1):
        cell = scores_sheet.cell(1, col_idx, header)
        cell.font = bold
        if header in metrics:
            if "BUETCSE" in header:
                cell.fill = buetcse_fill
            elif "SagorSarkar" in header:
                cell.fill = sagorsarkar_fill
            elif header in TRADITIONAL_METRICS:
                cell.fill = traditional_fill
            else:
                cell.fill = header_fill
            metric_columns[header] = get_column_letter(col_idx)
        else:
            cell.fill = header_fill

    for row_idx, (_, row) in enumerate(labeled_scores.iterrows(), start=2):
        for col_idx, header in enumerate(score_headers, start=1):
            scores_sheet.cell(row_idx, col_idx, excel_safe(row[header]))

    last_row = len(labeled_scores) + 1
    label_col = get_column_letter(score_headers.index("Label") + 1)

    for index, metric in enumerate(metrics, start=5):
        if metric not in metric_columns:
            continue
        metric_col = metric_columns[metric]
        cell_metric = dashboard.cell(index, 1, metric)

        if "BUETCSE" in metric:
            cell_metric.fill = buetcse_fill
        elif "SagorSarkar" in metric:
            cell_metric.fill = sagorsarkar_fill
        elif metric in TRADITIONAL_METRICS:
            cell_metric.fill = traditional_fill

        tp = f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},">="&$B$1,scores!${label_col}$2:${label_col}${last_row},1)'
        fp = f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},">="&$B$1,scores!${label_col}$2:${label_col}${last_row},0)'
        tn = f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},"<"&$B$1,scores!${label_col}$2:${label_col}${last_row},0)'
        fn = f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},"<"&$B$1,scores!${label_col}$2:${label_col}${last_row},1)'

        dashboard.cell(index, 2, tp)
        dashboard.cell(index, 3, fp)
        dashboard.cell(index, 4, tn)
        dashboard.cell(index, 5, fn)
        dashboard.cell(index, 6, f"=(B{index}+D{index})/SUM(B{index}:E{index})")
        dashboard.cell(index, 7, f"=IFERROR(B{index}/(B{index}+C{index}),0)")
        dashboard.cell(index, 8, f"=IFERROR(B{index}/(B{index}+E{index}),0)")
        dashboard.cell(index, 9, f"=IFERROR(2*G{index}*H{index}/(G{index}+H{index}),0)")

        for col_idx in range(6, 10):
            dashboard.cell(index, col_idx).number_format = "0.0000"

    metric_width = max(28, min(42, max((len(m) for m in metrics), default=16) + 2))
    dashboard.column_dimensions["A"].width = metric_width
    for col in ["B", "C", "D", "E"]:
        dashboard.column_dimensions[col].width = 10
    for col in ["F", "G", "H", "I"]:
        dashboard.column_dimensions[col].width = 12
    dashboard.column_dimensions["K"].width = 42
    for col in ["L", "M", "N"]:
        dashboard.column_dimensions[col].width = 24

    for col_idx, header in enumerate(score_headers, start=1):
        width = 14
        if header in {"text1", "text2", "text1_raw", "text2_raw"}:
            width = 50
        elif header in {"variant_type", "pair_type", "preprocess_profile", "preprocess_steps",
                        "source_link_text1", "source_link_text2"}:
            width = 22
        scores_sheet.column_dimensions[get_column_letter(col_idx)].width = width

    scores_sheet.freeze_panes = "A2"
    dashboard.freeze_panes = "A5"

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)
    print(f"  Interactive workbook: {output_file}")


# ═══════════════════════════════════════════════════════════════════
# 3. Pipeline
# ═══════════════════════════════════════════════════════════════════

def run_evaluation_for_dataset(
    *,
    input_path: Path,
    output_dir: Path,
    dataset_label: str,
    batch_size: int,
    max_length: int,
    device: str,
    threshold: float,
    sweep_thresholds: list[float],
) -> dict[str, object]:
    banner(f"BanglaBERT - {dataset_label}")
    print(f"  Input : {input_path}")
    print(f"  Output: {output_dir}")

    df = pd.read_excel(input_path, engine="openpyxl")
    print(f"  Loaded {len(df)} pairs from {input_path.name}")

    reference_ids = df["pair_no"].tolist()
    reference_labels = df["Label"].tolist()

    for col in ("text1", "text2"):
        df[col] = df[col].fillna("").astype(str)

    traditional_available = [m for m in TRADITIONAL_METRICS if m in df.columns]
    print(f"  Traditional metrics carried over: {traditional_available}")
    if len(traditional_available) != len(TRADITIONAL_METRICS):
        raise RuntimeError(f"Missing traditional metrics in {input_path.name}")

    all_vectors: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    model_stats: list[dict[str, object]] = []

    for model_name, cosine_column in MODELS:
        print(f"\n  --- Model: {model_name} ---")
        set_global_seed(SEED)
        df, t1_vecs, t2_vecs, stats = compute_banglabert_scores(
            df,
            model_name=model_name,
            cosine_column=cosine_column,
            batch_size=batch_size,
            max_length=max_length,
            device=device,
        )
        all_vectors[cosine_column] = (t1_vecs, t2_vecs)
        model_stats.append(stats)
        print(f"  [OK] {cosine_column}: mean={stats['cosine_mean']}, "
              f"range=[{stats['cosine_min']}, {stats['cosine_max']}], dim={stats['embedding_dim']}, "
              f"truncated={stats['truncated_at_max_length']}/{stats['unique_texts']}")

    # --- integrity guards -------------------------------------------------
    if df["pair_no"].tolist() != reference_ids:
        raise RuntimeError("pair_no order changed during BanglaBERT scoring.")
    if df["Label"].tolist() != reference_labels:
        raise RuntimeError("Label values changed during BanglaBERT scoring.")

    positive_count = int(df["Label"].sum())
    print(f"\n  Labels: {positive_count} positive, {len(df) - positive_count} negative (total {len(df)})")

    bert_metrics = [col for _, col in MODELS]
    all_metrics = bert_metrics + traditional_available
    print(f"  Metrics to evaluate: {all_metrics}")

    output_dir.mkdir(parents=True, exist_ok=True)

    # --- per-model score workbooks + embedding archives -------------------
    for model_name, cosine_column in MODELS:
        model_short = "buetcse" if "BUETCSE" in cosine_column else "sagorsarkar"
        scores_xlsx = output_dir / f"{model_short}_banglabert_scores.xlsx"
        vectors_npz = output_dir / f"{model_short}_banglabert_vectors.npz"

        t1_vecs, t2_vecs = all_vectors[cosine_column]
        score_df = df.copy()
        score_df["BanglaBERT_Model"] = model_name
        score_df["BanglaBERT_Pooling"] = "mean_pooling_l2_normalized"
        score_df["BanglaBERT_Dim"] = t1_vecs.shape[1] if len(t1_vecs) > 0 else 0
        score_df.to_excel(scores_xlsx, index=False, engine="openpyxl")
        print(f"  Scores workbook: {scores_xlsx}")

        np.savez_compressed(
            vectors_npz,
            pair_no=df["pair_no"].to_numpy(),
            text1_embeddings=t1_vecs,
            text2_embeddings=t2_vecs,
        )
        print(f"  Vectors archive: {vectors_npz}")

    # --- evaluation -------------------------------------------------------
    class_df = classification_metrics(df, threshold, all_metrics)
    sweep_df = threshold_sweep(df, sweep_thresholds, all_metrics)
    best_df = best_thresholds(sweep_df)

    class_df.to_csv(output_dir / "classification_metrics_threshold_0_50.csv", index=False, encoding="utf-8-sig")
    sweep_df.to_csv(output_dir / "threshold_sweep.csv", index=False, encoding="utf-8-sig")
    best_df.to_csv(output_dir / "best_threshold_metrics.csv", index=False, encoding="utf-8-sig")
    df.to_csv(output_dir / "labeled_scores.csv", index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(output_dir / "evaluation_results.xlsx", engine="openpyxl") as writer:
        class_df.to_excel(writer, sheet_name="threshold_0_50", index=False)
        sweep_df.to_excel(writer, sheet_name="threshold_sweep", index=False)
        best_df.to_excel(writer, sheet_name="best_thresholds", index=False)
        df.to_excel(writer, sheet_name="labeled_scores", index=False)

    build_interactive_workbook(
        df,
        all_metrics,
        output_dir / "interactive_threshold_evaluation.xlsx",
        source_description=(
            f"Dataset: {dataset_label} | Models: csebuetnlp/banglabert + sagorsarker/bangla-bert-base"
        ),
    )

    pd.DataFrame(model_stats).to_csv(
        output_dir / "encoder_diagnostics.csv", index=False, encoding="utf-8-sig"
    )

    print(f"\n  --- Classification metrics @ t={threshold} ({dataset_label}) ---")
    print(class_df.to_string(index=False))
    print(f"\n  --- Best threshold per metric ({dataset_label}) ---")
    print(best_df.to_string(index=False))
    print(f"  [OK] Outputs written to: {output_dir}\n")

    return {
        "dataset_label": dataset_label,
        "input": str(input_path),
        "output_dir": str(output_dir),
        "rows": int(len(df)),
        "positives": positive_count,
        "negatives": int(len(df) - positive_count),
        "metrics": all_metrics,
        "encoder_diagnostics": model_stats,
        "batch_size": batch_size,
        "max_length": max_length,
        "device": device,
        "pooling": "mean_pooling_l2_normalized",
        "seed": SEED,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="BanglaBERT evaluation for BanglaPlagPair-1K (2 models x raw/polished)."
    )
    parser.add_argument("--run", choices=["both", "normal", "polished"], default="both")
    parser.add_argument("--normal-input", type=Path, default=RAW_SCORES_XLSX)
    parser.add_argument("--polished-input", type=Path, default=POLISHED_SCORES_XLSX)
    parser.add_argument("--output-base", type=Path, default=BANGLABERT_DIR)
    parser.add_argument("--batch-size", type=int, default=8)
    parser.add_argument("--max-length", type=int, default=512)
    parser.add_argument("--device", default="cpu")
    parser.add_argument("--threshold", type=float, default=0.50)
    parser.add_argument(
        "--sweep",
        default=",".join(f"{t:g}" for t in THRESHOLDS),
        help="Comma-separated thresholds for sensitivity analysis",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    set_global_seed(SEED)
    ensure_dirs()
    sweep_thresholds = [float(v.strip()) for v in args.sweep.split(",") if v.strip()]

    common = dict(
        batch_size=args.batch_size,
        max_length=args.max_length,
        device=args.device,
        threshold=args.threshold,
        sweep_thresholds=sweep_thresholds,
    )

    reports: list[dict[str, object]] = []

    if args.run in ("both", "normal"):
        reports.append(
            run_evaluation_for_dataset(
                input_path=args.normal_input,
                output_dir=args.output_base / "normal",
                dataset_label="Normal (Raw) BanglaPlagPair-1K",
                **common,
            )
        )

    if args.run in ("both", "polished"):
        reports.append(
            run_evaluation_for_dataset(
                input_path=args.polished_input,
                output_dir=args.output_base / "polished",
                dataset_label="Polished (Preprocessed) BanglaPlagPair-1K",
                **common,
            )
        )

    (VALIDATION_DIR / "step4_banglabert_validation.json").write_text(
        json.dumps(reports, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    banner("STEP 4 COMPLETE")


if __name__ == "__main__":
    main()
