from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent


def excel_safe(value: object) -> object:
    if pd.isna(value):
        return ""
    return value


def resolve_metrics(df: pd.DataFrame, metrics_arg: str | None) -> list[str]:
    if metrics_arg:
        metrics = [value.strip() for value in metrics_arg.split(",") if value.strip()]
        missing = [metric for metric in metrics if metric not in df.columns]
        if missing:
            raise ValueError(f"Missing metric columns: {missing}")
        return metrics

    excluded = {
        "model",
        "pair_no",
        "dataset",
        "source_case_no",
        "variant_type",
        "text1",
        "text2",
        "text1_raw",
        "text2_raw",
        "preprocess_profile",
        "preprocess_steps",
        "text1_token_count",
        "text2_token_count",
        "model_detection_category",
        "Label",
    }
    metrics: list[str] = []
    for column in df.columns:
        if column in excluded:
            continue
        if pd.api.types.is_numeric_dtype(df[column]):
            metrics.append(column)
    if not metrics:
        raise ValueError("No numeric metric columns found for interactive workbook.")
    return metrics


def build_workbook(input_file: Path, output_file: Path, metrics_arg: str | None = None) -> None:
    df = pd.read_excel(input_file, engine="openpyxl")
    if "Label" not in df.columns:
        if "dataset" in df.columns:
            df["Label"] = df["dataset"].astype(str).str.lower().ne("mismatch").astype(int)
        elif "variant_type" in df.columns:
            df["Label"] = df["variant_type"].astype(str).str.lower().ne("mismatch").astype(int)
        else:
            raise ValueError("Input file must contain either Label, dataset, or variant_type.")

    metrics = resolve_metrics(df, metrics_arg)

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "dashboard"
    scores = wb.create_sheet("scores")

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    metric_fill = PatternFill(fill_type="solid", fgColor="FCE5CD")
    bold = Font(bold=True)

    dashboard["A1"] = "Threshold"
    dashboard["A1"].font = bold
    dashboard["B1"] = 0.5
    dashboard["B1"].number_format = "0.00"
    dashboard["K1"] = "Change B1 manually to observe result changes"
    dashboard["K1"].font = Font(italic=True)
    dashboard["K2"] = f"Source file: {input_file.name}"
    dashboard["K2"].font = Font(italic=True)

    header_row = 4
    for column, label in enumerate(["Metric", "TP", "FP", "TN", "FN", "Accuracy", "Precision", "Recall", "F1"], start=1):
        dashboard.cell(header_row, column, label).font = bold
        dashboard.cell(header_row, column).fill = header_fill

    score_headers = df.columns.tolist()
    metric_columns: dict[str, str] = {}

    for col_idx, header in enumerate(score_headers, start=1):
        cell = scores.cell(1, col_idx, header)
        cell.font = bold
        cell.fill = metric_fill if header in metrics else header_fill
        if header in metrics:
            metric_columns[header] = get_column_letter(col_idx)

    for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
        for col_idx, header in enumerate(score_headers, start=1):
            scores.cell(row_idx, col_idx, excel_safe(row[header]))

    last_row = len(df) + 1
    label_col = get_column_letter(score_headers.index("Label") + 1)

    for index, metric in enumerate(metrics, start=5):
        metric_col = metric_columns[metric]
        dashboard.cell(index, 1, metric)

        dashboard.cell(index, 2, f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},">="&$B$1,scores!${label_col}$2:${label_col}${last_row},1)')
        dashboard.cell(index, 3, f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},">="&$B$1,scores!${label_col}$2:${label_col}${last_row},0)')
        dashboard.cell(index, 4, f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},"<"&$B$1,scores!${label_col}$2:${label_col}${last_row},0)')
        dashboard.cell(index, 5, f'=COUNTIFS(scores!${metric_col}$2:${metric_col}${last_row},"<"&$B$1,scores!${label_col}$2:${label_col}${last_row},1)')
        dashboard.cell(index, 6, f"=(B{index}+D{index})/SUM(B{index}:E{index})")
        dashboard.cell(index, 7, f'=IFERROR(B{index}/(B{index}+C{index}),0)')
        dashboard.cell(index, 8, f'=IFERROR(B{index}/(B{index}+E{index}),0)')
        dashboard.cell(index, 9, f'=IFERROR(2*G{index}*H{index}/(G{index}+H{index}),0)')

        for col_idx in range(6, 10):
            dashboard.cell(index, col_idx).number_format = "0.0000"

    metric_width = max(16, min(42, max(len(metric) for metric in metrics) + 2))
    dashboard.column_dimensions["A"].width = metric_width
    for col in ["B", "C", "D", "E"]:
        dashboard.column_dimensions[col].width = 10
    for col in ["F", "G", "H", "I"]:
        dashboard.column_dimensions[col].width = 12
    dashboard.column_dimensions["K"].width = 42

    for col_idx, header in enumerate(score_headers, start=1):
        width = 14
        if header in {"text1", "text2"}:
            width = 50
        elif header in {"dataset", "variant_type", "model"}:
            width = 18
        elif header == "model_detection_category":
            width = 30
        scores.column_dimensions[get_column_letter(col_idx)].width = width

    scores.freeze_panes = "A2"
    dashboard.freeze_panes = "A4"

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_file)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build interactive threshold evaluation workbook.")
    parser.add_argument("--input", type=Path, required=True, help="Scored workbook with dataset and metric columns")
    parser.add_argument("--output", type=Path, required=True, help="Output interactive workbook path")
    parser.add_argument("--metrics", default=None, help="Comma-separated metric column names to include")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    build_workbook(args.input, args.output, args.metrics)
    print(f"Interactive workbook written to: {args.output}")

