# -*- coding: utf-8 -*-
"""Independent integrity audit for the paid-LLM evaluation artefacts."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
WORK_ROOT = SCRIPT_DIR.parent
PAID_DIR = WORK_ROOT / "paid llm evualation"
BASE_PATH = WORK_ROOT / "dataset" / "banglaplagpair1k_base.xlsx"
SCORES_PATH = PAID_DIR / "combined_1000_testcases.xlsx"
OUT_PATH = WORK_ROOT / "validation" / "step10_paid_llm_validation.json"

EXPECTED_SCORE_SHA256 = "34e4d34fa0214c55567939d9d1c404241253fdac29df058bc49b61fdf1466097"
THRESHOLDS = [round(value, 2) for value in np.arange(0.10, 0.91, 0.05)]
FIXED_THRESHOLD = 0.50
METRIC_COLUMNS = [
    "Metric", "Threshold", "Accuracy", "Precision", "Recall", "F1",
    "TP", "FP", "TN", "FN",
]
SUMMARY_COLUMNS = [
    "Model", "Metric", "Threshold_0_50_Accuracy", "Threshold_0_50_Precision",
    "Threshold_0_50_Recall", "Threshold_0_50_F1", "Best_Threshold",
    "Best_Accuracy", "Best_Precision", "Best_Recall", "Best_F1", "Output_Folder",
]
EXPECTED_MODEL_FILES = {
    "normalized_scores.xlsx",
    "labeled_scores.csv",
    "classification_metrics_threshold_0_50.csv",
    "threshold_sweep.csv",
    "best_threshold_metrics.csv",
    "evaluation_results.xlsx",
}

MODELS = [
    ("DeepSeek V4 Flash", "deepseek_v4_flash", "DeepSeek_V4_Flash_Score"),
    ("GPT 5.6 Luna", "gpt_5_6_luna", "GPT_5_6_Luna_Score"),
    (
        "Gemini 3.5 Flash Lite",
        "gemini_3_5_flash_lite",
        "Gemini_3_5_Flash_Lite_Score",
    ),
    ("Qwen 3.7 Flash", "qwen_3_7_flash", "Qwen_3_7_Flash_Score"),
    ("Claude Sonnet 5", "claude_sonnet_5", "Claude_Sonnet_5_Score"),
    ("Gemini 3.6 Flash", "gemini_3_6_flash", "Gemini_3_6_Flash_Score"),
    (
        "DeepSeek V4 Pro 0813",
        "deepseek_v4_pro_0813",
        "DeepSeek_V4_Pro_0813_Score",
    ),
    ("Grok 4.3", "grok_4_3", "Grok_4_3_Score"),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def metrics_for(labels: pd.Series, scores: pd.Series, threshold: float, metric: str) -> dict:
    predicted = scores.ge(threshold).astype(int)
    actual = labels.astype(int)
    tp = int(((predicted == 1) & (actual == 1)).sum())
    fp = int(((predicted == 1) & (actual == 0)).sum())
    tn = int(((predicted == 0) & (actual == 0)).sum())
    fn = int(((predicted == 0) & (actual == 1)).sum())
    accuracy = (tp + tn) / len(actual)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {
        "Metric": metric,
        "Threshold": threshold,
        "Accuracy": round(accuracy, 6),
        "Precision": round(precision, 6),
        "Recall": round(recall, 6),
        "F1": round(f1, 6),
        "TP": tp,
        "FP": fp,
        "TN": tn,
        "FN": fn,
    }


def expected_tables(labels: pd.Series, scores: pd.Series, metric: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    fixed = pd.DataFrame([metrics_for(labels, scores, FIXED_THRESHOLD, metric)])
    sweep = pd.DataFrame([metrics_for(labels, scores, threshold, metric) for threshold in THRESHOLDS])
    best = sweep.sort_values(
        ["Metric", "F1", "Accuracy", "Recall", "Precision"],
        ascending=[True, False, False, False, False],
    ).groupby("Metric", as_index=False).head(1).reset_index(drop=True)
    return fixed, sweep, best


def assert_frame(actual: pd.DataFrame, expected: pd.DataFrame, name: str) -> None:
    actual = actual.reset_index(drop=True)
    expected = expected.reset_index(drop=True)
    if list(actual.columns) != list(expected.columns):
        raise AssertionError(f"{name}: columns {list(actual.columns)} != {list(expected.columns)}")
    pd.testing.assert_frame_equal(actual, expected, check_dtype=False, atol=1e-12, rtol=0)


def workbook_audit(directory: Path, slug: str, metric_count: int, row_count: int) -> dict:
    interactive_name = (
        "interactive_threshold_combined_llm_models.xlsx"
        if slug == "combined_llm_models"
        else f"interactive_threshold_{slug}.xlsx"
    )
    expected_files = EXPECTED_MODEL_FILES | {interactive_name}
    actual_files = {path.name for path in directory.iterdir() if path.is_file()}
    if actual_files != expected_files:
        raise AssertionError(
            f"{directory.name}: file set mismatch; missing={expected_files-actual_files}, "
            f"extra={actual_files-expected_files}"
        )

    normalized = pd.ExcelFile(directory / "normalized_scores.xlsx", engine="openpyxl")
    if normalized.sheet_names != ["Sheet1"]:
        raise AssertionError(f"{directory.name}: unexpected normalized workbook sheets")

    results = pd.ExcelFile(directory / "evaluation_results.xlsx", engine="openpyxl")
    expected_result_sheets = ["threshold_0_50", "threshold_sweep", "best_thresholds", "labeled_scores"]
    if results.sheet_names != expected_result_sheets:
        raise AssertionError(f"{directory.name}: unexpected evaluation workbook sheets")

    workbook = load_workbook(directory / interactive_name, data_only=False, read_only=False)
    if workbook.sheetnames != ["dashboard", "scores"]:
        raise AssertionError(f"{directory.name}: unexpected interactive workbook sheets")
    dashboard = workbook["dashboard"]
    score_sheet = workbook["scores"]
    if dashboard["A1"].value != "Threshold" or float(dashboard["B1"].value) != FIXED_THRESHOLD:
        raise AssertionError(f"{directory.name}: interactive threshold control is invalid")
    if score_sheet.max_row != row_count + 1:
        raise AssertionError(f"{directory.name}: interactive score row count is invalid")

    formula_cells = []
    formula_errors = []
    for row in dashboard.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formula_cells.append(cell.coordinate)
                if any(token in value for token in ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")):
                    formula_errors.append(cell.coordinate)
    if len(formula_cells) != metric_count * 8 or formula_errors:
        raise AssertionError(
            f"{directory.name}: formula audit failed ({len(formula_cells)} formulas, errors={formula_errors})"
        )
    workbook.close()

    return {
        "files": sorted(actual_files),
        "result_sheets": expected_result_sheets,
        "interactive_sheets": ["dashboard", "scores"],
        "interactive_formula_count": len(formula_cells),
        "formula_errors": formula_errors,
    }


def score_distribution(labels: pd.Series, scores: pd.Series) -> dict:
    positive = scores[labels.eq(1)]
    negative = scores[labels.eq(0)]
    return {
        "distinct_scores": int(scores.nunique()),
        "positive_mean": round(float(positive.mean()), 6),
        "positive_min": round(float(positive.min()), 6),
        "positive_max": round(float(positive.max()), 6),
        "negative_mean": round(float(negative.mean()), 6),
        "negative_min": round(float(negative.min()), 6),
        "negative_max": round(float(negative.max()), 6),
        "class_separation_margin": round(float(positive.min() - negative.max()), 6),
    }


def main() -> None:
    score_hash = sha256(SCORES_PATH)
    if score_hash != EXPECTED_SCORE_SHA256:
        raise AssertionError(
            f"Paid-LLM source workbook hash changed: {score_hash} != {EXPECTED_SCORE_SHA256}"
        )

    base = pd.read_excel(BASE_PATH, engine="openpyxl")
    source = pd.read_excel(SCORES_PATH, engine="openpyxl")
    if len(base) != 1000 or len(source) != 1000:
        raise AssertionError("Expected exactly 1,000 canonical and paid-score rows.")
    if not base["pair_id"].is_unique or not source["pair_id"].is_unique:
        raise AssertionError("pair_id must be unique in both inputs.")

    canonical_label = base["pair_type"].astype(str).str.lower().ne("mismatch").astype(int)
    if not canonical_label.equals(base["Label"].astype(int)):
        raise AssertionError("Canonical labels disagree with pair_type.")
    if not canonical_label.equals(base["binary_label"].astype(int)):
        raise AssertionError("Canonical labels disagree with binary_label.")

    joined = base[["pair_id", "pair_no", "text1", "text2", "variant_type", "Label"]].merge(
        source[["pair_id", "text1", "text2"] + [metric for _, _, metric in MODELS]],
        on="pair_id",
        validate="one_to_one",
        suffixes=("_base", "_paid"),
    )
    if not joined["pair_id"].tolist() == list(range(1, 1001)):
        raise AssertionError("pair_id order is not 1..1000.")
    if not joined["text1_base"].astype(str).equals(joined["text1_paid"].astype(str)):
        raise AssertionError("Source text1 differs from the canonical benchmark.")
    if not joined["text2_base"].astype(str).equals(joined["text2_paid"].astype(str)):
        raise AssertionError("Source text2 differs from the canonical benchmark.")
    labels = joined["Label"].astype(int)

    report: dict[str, object] = {
        "status": "PASS",
        "method": {
            "fixed_threshold": FIXED_THRESHOLD,
            "thresholds": THRESHOLDS,
            "positive_rule": "score >= threshold",
            "best_threshold_tiebreak": ["F1", "Accuracy", "Recall", "Precision"],
            "metric_columns": METRIC_COLUMNS,
        },
        "inputs": {
            "score_workbook": str(SCORES_PATH),
            "score_workbook_sha256": score_hash,
            "canonical_base": str(BASE_PATH),
            "canonical_base_sha256": sha256(BASE_PATH),
            "rows": len(joined),
            "positive": int(labels.sum()),
            "negative": int((1 - labels).sum()),
            "pair_id_order": "1..1000",
            "text1_exact_matches": int(joined["text1_base"].astype(str).eq(joined["text1_paid"].astype(str)).sum()),
            "text2_exact_matches": int(joined["text2_base"].astype(str).eq(joined["text2_paid"].astype(str)).sum()),
        },
        "models": {},
    }

    all_expected_fixed = []
    all_expected_sweep = []
    all_expected_best = []
    expected_summary_rows = []

    for model_name, slug, metric in MODELS:
        numeric_scores = pd.to_numeric(joined[metric], errors="coerce")
        if numeric_scores.isna().any() or not numeric_scores.between(0, 1).all():
            raise AssertionError(f"{metric}: invalid score values")
        expected_fixed, expected_sweep, expected_best = expected_tables(labels, numeric_scores, metric)
        directory = PAID_DIR / f"{slug}_evaluation"

        actual_labeled = pd.read_csv(directory / "labeled_scores.csv")
        if len(actual_labeled) != 1000:
            raise AssertionError(f"{metric}: labeled score row count differs")
        if not actual_labeled["pair_no"].tolist() == list(range(1, 1001)):
            raise AssertionError(f"{metric}: pair_no order differs")
        if not actual_labeled["Label"].astype(int).equals(labels):
            raise AssertionError(f"{metric}: labels differ")
        if not np.allclose(actual_labeled[metric].astype(float), numeric_scores, atol=0, rtol=0):
            raise AssertionError(f"{metric}: emitted scores differ from source")

        actual_fixed = pd.read_csv(directory / "classification_metrics_threshold_0_50.csv")
        actual_sweep = pd.read_csv(directory / "threshold_sweep.csv")
        actual_best = pd.read_csv(directory / "best_threshold_metrics.csv")
        assert_frame(actual_fixed, expected_fixed, f"{metric} fixed")
        assert_frame(actual_sweep, expected_sweep, f"{metric} sweep")
        assert_frame(actual_best, expected_best, f"{metric} best")

        with pd.ExcelFile(directory / "evaluation_results.xlsx", engine="openpyxl") as workbook:
            assert_frame(pd.read_excel(workbook, "threshold_0_50"), expected_fixed, f"{metric} xlsx fixed")
            assert_frame(pd.read_excel(workbook, "threshold_sweep"), expected_sweep, f"{metric} xlsx sweep")
            assert_frame(pd.read_excel(workbook, "best_thresholds"), expected_best, f"{metric} xlsx best")

        workbook_checks = workbook_audit(directory, slug, metric_count=1, row_count=1000)
        fixed_row = expected_fixed.iloc[0].to_dict()
        best_row = expected_best.iloc[0].to_dict()
        fixed_row["Specificity"] = round(fixed_row["TN"] / (fixed_row["TN"] + fixed_row["FP"]), 6)
        best_row["Specificity"] = round(best_row["TN"] / (best_row["TN"] + best_row["FP"]), 6)
        report["models"][metric] = {
            "model": model_name,
            "fixed": fixed_row,
            "best": best_row,
            "distribution": score_distribution(labels, numeric_scores),
            "artefacts": workbook_checks,
        }

        expected_summary_rows.append(
            {
                "Model": model_name,
                "Metric": metric,
                "Threshold_0_50_Accuracy": fixed_row["Accuracy"],
                "Threshold_0_50_Precision": fixed_row["Precision"],
                "Threshold_0_50_Recall": fixed_row["Recall"],
                "Threshold_0_50_F1": fixed_row["F1"],
                "Best_Threshold": best_row["Threshold"],
                "Best_Accuracy": best_row["Accuracy"],
                "Best_Precision": best_row["Precision"],
                "Best_Recall": best_row["Recall"],
                "Best_F1": best_row["F1"],
                "Output_Folder": str(directory.relative_to(WORK_ROOT)),
            }
        )
        all_expected_fixed.append(expected_fixed)
        all_expected_sweep.append(expected_sweep)
        all_expected_best.append(expected_best)

    combined_dir = PAID_DIR / "combined_llm_model_evaluation"
    combined_fixed = pd.concat(all_expected_fixed, ignore_index=True)
    combined_sweep = pd.concat(
        [
            pd.concat(
                [frame.loc[frame["Threshold"].eq(threshold)] for frame in all_expected_sweep],
                ignore_index=True,
            )
            for threshold in THRESHOLDS
        ],
        ignore_index=True,
    )
    combined_best = pd.concat(all_expected_best, ignore_index=True).sort_values("Metric").reset_index(drop=True)

    assert_frame(
        pd.read_csv(combined_dir / "classification_metrics_threshold_0_50.csv"),
        combined_fixed,
        "combined fixed",
    )
    assert_frame(pd.read_csv(combined_dir / "threshold_sweep.csv"), combined_sweep, "combined sweep")
    actual_combined_best = pd.read_csv(combined_dir / "best_threshold_metrics.csv").sort_values("Metric").reset_index(drop=True)
    assert_frame(actual_combined_best, combined_best, "combined best")
    combined_workbooks = workbook_audit(
        combined_dir,
        "combined_llm_models",
        metric_count=len(MODELS),
        row_count=1000,
    )

    summary_csv = pd.read_csv(PAID_DIR / "llm_model_evaluation_summary.csv")
    expected_summary = pd.DataFrame(expected_summary_rows)
    if list(summary_csv.columns) != SUMMARY_COLUMNS:
        raise AssertionError("Summary CSV does not match the historical 12-column schema.")
    assert_frame(summary_csv, expected_summary, "summary CSV")
    summary_xlsx = pd.read_excel(PAID_DIR / "llm_model_evaluation_summary.xlsx", engine="openpyxl")
    assert_frame(summary_xlsx, expected_summary, "summary XLSX")

    report["combined"] = {
        "models": len(MODELS),
        "fixed_rows": len(combined_fixed),
        "sweep_rows": len(combined_sweep),
        "best_rows": len(combined_best),
        "artefacts": combined_workbooks,
    }
    report["checks"] = {
        "source_hash_unchanged": True,
        "score_columns_unchanged": True,
        "labels_match_pair_type_and_binary_label": True,
        "pair_ids_and_texts_match_canonical_base": True,
        "csv_metrics_match_independent_recalculation": True,
        "xlsx_metric_sheets_match_independent_recalculation": True,
        "summary_csv_xlsx_agree": True,
        "interactive_workbook_formulas_valid": True,
        "historical_file_and_table_schema_preserved": True,
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PASS: paid-LLM evaluation validated; report written to {OUT_PATH}")


if __name__ == "__main__":
    main()
