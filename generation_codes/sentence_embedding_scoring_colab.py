# =====================================================
# Sentence Embedding Similarity Evaluation for Excel
# Colab-ready Code
# =====================================================

# In a Colab cell, install first with:
# !pip install -q sentence-transformers openpyxl

import os
import gc
from pathlib import Path

import torch
from openpyxl import load_workbook
from sentence_transformers import SentenceTransformer, util

# =====================================================
# Configuration
# =====================================================

WORKBOOK_PATH = "./data/input.xlsx"  # Set this to your dataset workbook path

# =====================================================
# Select Model
# Keep only one MODEL_NAME uncommented
# =====================================================

# Supported Model IDs:
# 1) sentence-transformers/LaBSE
# 2) intfloat/multilingual-e5-base
# 3) BAAI/bge-m3
# 4) sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2

# Lightweight multilingual SBERT baseline
MODEL_NAME = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
SCORE_COLUMN_NAME = "MiniLM_Multilingual_Score"


# =====================================================
# Helper Functions
# =====================================================

def get_column_index(worksheet, header_name: str) -> int:
    expected = header_name.strip().lower()

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value
        if isinstance(value, str) and value.strip().lower() == expected:
            return index

    raise ValueError(f"Column '{header_name}' was not found in the workbook header.")


def get_or_create_column_index(worksheet, header_name: str) -> int:
    expected = header_name.strip().lower()
    blank_column = None

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value

        if isinstance(value, str) and value.strip().lower() == expected:
            return index

        if blank_column is None and (value is None or (isinstance(value, str) and not value.strip())):
            blank_column = index

    if blank_column is None:
        blank_column = worksheet.max_column + 1

    worksheet.cell(row=1, column=blank_column).value = header_name
    return blank_column


def should_generate_score(existing_score) -> bool:
    if existing_score is None:
        return True

    if isinstance(existing_score, str) and not existing_score.strip():
        return True

    return False


def prepare_texts_for_model(text1: str, text2: str):
    """
    E5 models require the 'query:' prefix.
    """
    if "e5" in MODEL_NAME.lower():
        return [
            "query: " + text1,
            "query: " + text2
        ]

    return [text1, text2]


def load_embedding_model(model_name: str):
    print(f"Loading model: {model_name}")

    device = "cuda" if torch.cuda.is_available() else "cpu"
    print("Device:", device)

    model = SentenceTransformer(model_name, device=device)
    return model, device


def calculate_similarity_score(model, text1: str, text2: str) -> float:
    texts = prepare_texts_for_model(text1, text2)

    embeddings = model.encode(
        texts,
        convert_to_tensor=True,
        normalize_embeddings=True,
        show_progress_bar=False
    )

    score = util.cos_sim(embeddings[0], embeddings[1]).item()

    return float(score)


def process_workbook(workbook_path: str, model, score_column_name: str) -> None:
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(f"File not found: {workbook_path}")

    print(f"Processing workbook: {workbook_path}")

    workbook = load_workbook(workbook_path)
    worksheet = workbook.active

    text1_column = get_column_index(worksheet, "text1")
    text2_column = get_column_index(worksheet, "text2")
    score_column = get_or_create_column_index(worksheet, score_column_name)

    total_rows = worksheet.max_row - 1

    for row_number in range(2, worksheet.max_row + 1):
        existing_score = worksheet.cell(row=row_number, column=score_column).value

        if not should_generate_score(existing_score):
            print(f"Skipping row {row_number - 1}/{total_rows}: existing score {existing_score}")
            continue

        text1 = worksheet.cell(row=row_number, column=text1_column).value
        text2 = worksheet.cell(row=row_number, column=text2_column).value

        if text1 is None or text2 is None:
            print(f"Skipping row {row_number - 1}/{total_rows}: missing text")
            continue

        try:
            score = calculate_similarity_score(
                model,
                str(text1),
                str(text2)
            )

            worksheet.cell(row=row_number, column=score_column).value = score

            # Save after each row
            workbook.save(workbook_path)

            print(f"Processed row {row_number - 1}/{total_rows}: {score:.4f}")

        except Exception as error:
            print(f"Error on row {row_number - 1}: {error}")

    workbook.save(workbook_path)
    print("Excel file updated successfully!")


def main():
    model, device = load_embedding_model(MODEL_NAME)

    try:
        process_workbook(
            WORKBOOK_PATH,
            model,
            SCORE_COLUMN_NAME
        )

        print("Done!")
        print("Model:", MODEL_NAME)
        print("Score column:", SCORE_COLUMN_NAME)

    finally:
        del model
        gc.collect()

        if torch.cuda.is_available():
            torch.cuda.empty_cache()


if __name__ == "__main__":
    main()
