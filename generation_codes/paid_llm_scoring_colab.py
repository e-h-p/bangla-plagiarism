import os
import re
import time

from openpyxl import load_workbook
from openai import OpenAI


# ============================================================
# 1. OPENROUTER API KEY
# ============================================================

OPENROUTER_API_KEY = os.environ.get("OPENROUTER_API_KEY")

if not OPENROUTER_API_KEY:
    raise RuntimeError(
        "OPENROUTER_API_KEY is not set. "
        "Set it as an environment variable before running the script."
    )


# ============================================================
# 2. OPENROUTER CLIENT
# ============================================================

client = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=OPENROUTER_API_KEY,
)


# ============================================================
# 3. DATASET PATH
# ============================================================

WORKBOOK_PATH = os.environ.get(
    "WORKBOOK_PATH",
    "./data/input.xlsx",
)


# ============================================================
# 4. PAID MODELS
# ============================================================

MODELS = {
    "DeepSeek_V4_Flash_Score": "deepseek/deepseek-v4-flash",
    "GPT_5_6_Luna_Score": "openai/gpt-5.6-luna",
    "Gemini_3_5_Flash_Lite_Score": "google/gemini-3.5-flash-lite",
    "Qwen_3_7_Flash_Score": "qwen/qwen3.7-flash",
    "Claude_Sonnet_5_Score": "anthropic/claude-sonnet-5",
    "Gemini_3_6_Flash_Score": "google/gemini-3.6-flash",
    "DeepSeek_V4_Pro_0813_Score": "deepseek/deepseek-v4-pro-0813",
    "Grok_4_3_Score": "x-ai/grok-4.3",
}


MAX_RETRIES = 5
REQUEST_DELAY_SECONDS = 0.2


# ============================================================
# 5. SHARED P0 PROMPT
# ============================================================

PROMPT_TEMPLATE = """Act as an expert Bangla NLP Researcher and Plagiarism Evaluator. Your task is to independently read and deeply understand the meaning between Text 1 and Text 2.

HUMAN-LIKE JUDGMENT: You must read the pair yourself, deeply analyze the underlying semantic meaning using your own cognition, and assign a similarity score manually based on the guidelines below. Do NOT write or rely on Python code (like TF-IDF, Cosine similarity, or BERT scripts) to evaluate or calculate the similarity.

GUIDELINES FOR BANGLA SEMANTICS:

- Focus entirely on the MEANING.
- Properly account for Bangla morphological variations, word-order changes, and lexical substitutions (e.g., Tatsama/Tadbhava synonyms like দেশ/রাষ্ট্র or চাঁদ/চন্দ্র). If the core meaning is preserved despite vocabulary changes, assign a high score.

CLASSIFICATION THRESHOLDS:

- 0.00 - 0.40: Independent / No Plagiarism
- 0.41 - 0.65: Suspicious / Light Paraphrasing
- 0.66 - 0.85: Heavy Paraphrasing / Probable Plagiarism
- 0.86 - 1.00: Direct Copy / Verbatim Plagiarism

Text 1: {text1}
Text 2: {text2}

Based on the guidelines, provide only the numerical Similarity Score (between 0.00 and 1.00). Do not include any other text, reasoning, or formatting."""


def build_prompt(text1: str, text2: str) -> str:
    return PROMPT_TEMPLATE.format(
        text1=text1,
        text2=text2,
    )


# ============================================================
# 6. SCORE EXTRACTION
# ============================================================

BANGLA_TO_ENGLISH_DIGITS = str.maketrans(
    "০১২৩৪৫৬৭৮৯",
    "0123456789",
)


def extract_score(output_text: str) -> float:
    normalized_text = str(output_text).translate(
        BANGLA_TO_ENGLISH_DIGITS
    )

    clean_text = re.sub(
        r"(text|পাঠ্য|ট্যাক্সট|টেক্সট)\s*[12]",
        "",
        normalized_text,
        flags=re.IGNORECASE,
    )

    matches = re.findall(
        r"(?<![\d.])(?:0(?:\.\d+)?|1(?:\.0+)?)(?![\d.])",
        clean_text,
    )

    if not matches:
        raise ValueError(
            f"No valid similarity score found: {output_text!r}"
        )

    score = float(matches[0])

    if not 0.0 <= score <= 1.0:
        raise ValueError(
            f"Score is outside the valid range: {score}"
        )

    return score


# ============================================================
# 7. EXCEL COLUMN FUNCTIONS
# ============================================================

def get_column_index(worksheet, header_name: str) -> int:
    expected = header_name.strip().lower()

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value

        if (
            isinstance(value, str)
            and value.strip().lower() == expected
        ):
            return index

    raise ValueError(
        f"Column '{header_name}' was not found in the workbook."
    )


def get_or_create_column_index(
    worksheet,
    header_name: str,
) -> int:
    expected = header_name.strip().lower()

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value

        if (
            isinstance(value, str)
            and value.strip().lower() == expected
        ):
            return index

    new_column = worksheet.max_column + 1

    worksheet.cell(
        row=1,
        column=new_column,
    ).value = header_name

    return new_column


def should_generate_score(existing_score) -> bool:
    if existing_score is None:
        return True

    if (
        isinstance(existing_score, str)
        and not existing_score.strip()
    ):
        return True

    return False


# ============================================================
# 8. API SCORE GENERATION
# ============================================================

def generate_score(
    model_id: str,
    text1: str,
    text2: str,
):
    prompt = build_prompt(text1, text2)
    last_error = None

    is_gemini_3 = model_id.startswith("google/gemini-3.")

    if is_gemini_3:
        reasoning_config = {
            "effort": "minimal",
            "exclude": True,
        }
        output_limit = 128
    else:
        reasoning_config = {
            "enabled": False,
        }
        output_limit = 12

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = client.chat.completions.create(
                model=model_id,
                messages=[
                    {
                        "role": "user",
                        "content": prompt,
                    }
                ],
                temperature=0,
                max_tokens=output_limit,
                extra_body={
                    "reasoning": reasoning_config
                },
            )

            response_text = (
                response.choices[0].message.content or ""
            ).strip()

            print(f"[Model Says]: {response_text}")

            score = extract_score(response_text)

            input_tokens = 0
            output_tokens = 0

            if response.usage is not None:
                input_tokens = (
                    response.usage.prompt_tokens or 0
                )
                output_tokens = (
                    response.usage.completion_tokens or 0
                )

            return score, input_tokens, output_tokens

        except Exception as error:
            last_error = error
            error_text = str(error)

            if "Error code: 400" in error_text:
                raise RuntimeError(
                    f"Invalid request configuration: {error}"
                )

            if attempt == MAX_RETRIES:
                break

            if (
                "Error code: 429" in error_text
                or "Rate limit" in error_text
            ):
                wait_seconds = 65
            else:
                wait_seconds = min(2 ** attempt, 30)

            print(
                f"API error: {error}\n"
                f"Retry {attempt}/{MAX_RETRIES} "
                f"after {wait_seconds} seconds..."
            )

            time.sleep(wait_seconds)

    raise RuntimeError(
        f"Request failed after {MAX_RETRIES} attempts: "
        f"{last_error}"
    )


# ============================================================
# 9. PROCESS ONE MODEL
# ============================================================

def process_model(
    workbook_path: str,
    model_id: str,
    score_column_name: str,
) -> None:
    print("\n" + "=" * 70)
    print(f"Model: {model_id}")
    print(f"Score column: {score_column_name}")
    print("=" * 70)

    workbook = load_workbook(workbook_path)
    worksheet = workbook.active

    text1_column = get_column_index(
        worksheet,
        "text1",
    )

    text2_column = get_column_index(
        worksheet,
        "text2",
    )

    score_column = get_or_create_column_index(
        worksheet,
        score_column_name,
    )

    workbook.save(workbook_path)

    total_rows = worksheet.max_row - 1
    processed_count = 0
    skipped_count = 0
    error_count = 0
    total_input_tokens = 0
    total_output_tokens = 0

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        dataset_row = row_number - 1

        existing_score = worksheet.cell(
            row=row_number,
            column=score_column,
        ).value

        if not should_generate_score(existing_score):
            skipped_count += 1

            print(
                f"Skipping row {dataset_row}/{total_rows}: "
                f"existing score = {existing_score}"
            )
            continue

        text1 = worksheet.cell(
            row=row_number,
            column=text1_column,
        ).value

        text2 = worksheet.cell(
            row=row_number,
            column=text2_column,
        ).value

        if (
            text1 is None
            or not str(text1).strip()
            or text2 is None
            or not str(text2).strip()
        ):
            error_count += 1

            print(
                f"Row {dataset_row}/{total_rows} skipped: "
                "text1 or text2 is blank."
            )
            continue

        try:
            score, input_tokens, output_tokens = generate_score(
                model_id=model_id,
                text1=str(text1),
                text2=str(text2),
            )

            worksheet.cell(
                row=row_number,
                column=score_column,
            ).value = score

            workbook.save(workbook_path)

            processed_count += 1
            total_input_tokens += input_tokens
            total_output_tokens += output_tokens

            print(
                f"Processed row {dataset_row}/{total_rows}: "
                f"score={score:.4f}, "
                f"input_tokens={input_tokens}, "
                f"output_tokens={output_tokens}"
            )

            time.sleep(
                REQUEST_DELAY_SECONDS
            )

        except Exception as error:
            error_count += 1

            print(
                f"Error on row "
                f"{dataset_row}/{total_rows}: "
                f"{error}"
            )

    workbook.save(workbook_path)
    workbook.close()

    print("\nModel processing complete.")
    print(f"Model: {model_id}")
    print(f"Processed: {processed_count}")
    print(f"Skipped existing: {skipped_count}")
    print(f"Errors or blank texts: {error_count}")
    print(f"Total input tokens: {total_input_tokens}")
    print(f"Total output tokens: {total_output_tokens}")


# ============================================================
# 10. MAIN
# ============================================================

def main():
    if not os.path.exists(WORKBOOK_PATH):
        raise FileNotFoundError(
            f"Excel file not found: {WORKBOOK_PATH}"
        )

    print(f"Dataset: {WORKBOOK_PATH}")
    print(f"Total models: {len(MODELS)}")

    for score_column_name, model_id in MODELS.items():
        try:
            process_model(
                workbook_path=WORKBOOK_PATH,
                model_id=model_id,
                score_column_name=score_column_name,
            )

        except Exception as error:
            print(
                f"\nModel processing failed: {model_id}\n"
                f"Error: {error}"
            )

    print("\nAll model processing completed.")
    print(f"Updated file: {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()