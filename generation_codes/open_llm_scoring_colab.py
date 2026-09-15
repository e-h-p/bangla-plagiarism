import os
import gc
import re

import torch
from openpyxl import load_workbook
from transformers import (
    AutoModelForCausalLM,
    AutoTokenizer,
    BitsAndBytesConfig,
)


# ============================================================
# 1. HUGGING FACE TOKEN
# ============================================================

HF_TOKEN = os.environ.get("HF_TOKEN")

if not HF_TOKEN:
    raise RuntimeError(
        "HF_TOKEN is not set in the runtime environment."
    )


# ============================================================
# 2. DATASET PATH
# ============================================================

WORKBOOK_PATH = "./data/input.xlsx"


# ============================================================
# 3. OPEN / OPEN-WEIGHT MODELS
#
# Run one model at a time.
# Uncomment one MODEL_ID + SCORE_COLUMN_NAME pair
# and keep the remaining pairs commented.
# ============================================================


# --- 1) Qwen2.5 0.5B Instruct ---

MODEL_ID = "Qwen/Qwen2.5-0.5B-Instruct"
SCORE_COLUMN_NAME = "Qwen2_5_0_5B_Score"


# --- 2) Qwen2.5 1.5B Instruct ---

# MODEL_ID = "Qwen/Qwen2.5-1.5B-Instruct"
# SCORE_COLUMN_NAME = "Qwen2.5-1.5B-Instruct"


# --- 3) Gemma 2 2B IT ---

# MODEL_ID = "google/gemma-2-2b-it"
# SCORE_COLUMN_NAME = "Gemma_2_2B_Score"


# --- 4) Gemma 2 9B IT ---

# MODEL_ID = "google/gemma-2-9b-it"
# SCORE_COLUMN_NAME = "Gemma_2_9B_Score"


# --- 5) Aya Expanse 8B ---

# MODEL_ID = "CohereForAI/aya-expanse-8b"
# SCORE_COLUMN_NAME = "Aya_Expanse_8B_Score"


# --- 6) Llama 3.2 3B Instruct ---

# MODEL_ID = "meta-llama/Llama-3.2-3B-Instruct"
# SCORE_COLUMN_NAME = "Llama_3_2_3B_Score"


# --- 7) Llama 3 8B Bangla 4bit ---

# MODEL_ID = "KillerShoaib/llama-3-8b-bangla-4bit"
# SCORE_COLUMN_NAME = "Llama_3_8B_Bangla_4bit_Score"


# --- 8) BanglaLLama 3 8B Instruct ---

# MODEL_ID = (
#     "BanglaLLM/"
#     "BanglaLLama-3-8b-bangla-alpaca-orca-instruct-v0.0.1"
# )
# SCORE_COLUMN_NAME = "BanglaLLama_3_8B_Score"


# ============================================================
# 4. SHARED P0 PROMPT
# ============================================================

PROMPT_TEMPLATE = """Act as an expert Bangla NLP Researcher and Plagiarism Evaluator. Your task is to independently read and deeply understand the meaning between Text 1 and Text 2.

HUMAN-LIKE JUDGMENT: You must read the pair yourself, deeply analyze the underlying semantic meaning using your own cognition, and assign a similarity score manually based on the guidelines below. Do NOT write or rely on Python code (like TF-IDF, Cosine similarity, or BERT scripts) to evaluate or calculate the similarity.

GUIDELINES FOR BANGLA SEMANTICS:

* Focus entirely on the MEANING.
* Properly account for Bangla morphological variations, word-order changes, and lexical substitutions (e.g., Tatsama/Tadbhava synonyms like দেশ/রাষ্ট্র or চাঁদ/চন্দ্র). If the core meaning is preserved despite vocabulary changes, assign a high score.

CLASSIFICATION THRESHOLDS:

* 0.00 - 0.40: Independent / No Plagiarism
* 0.41 - 0.65: Suspicious / Light Paraphrasing
* 0.66 - 0.85: Heavy Paraphrasing / Probable Plagiarism
* 0.86 - 1.00: Direct Copy / Verbatim Plagiarism

Text 1: {text1}
Text 2: {text2}

Based on the guidelines, provide only the numerical Similarity Score (between 0.00 and 1.00). Do not include any other text, reasoning, or formatting."""


def build_prompt(text1: str, text2: str) -> str:
    return PROMPT_TEMPLATE.format(
        text1=text1,
        text2=text2,
    )


# ============================================================
# 5. SCORE EXTRACTION
# ============================================================

def extract_score(output_text: str) -> float:

    # Remove Text 1 / Text 2 labels to avoid false extraction
    clean_text = re.sub(
        r"(text|পাঠ্য|ট্যাক্সট|টেক্সট)\s*[12১২]",
        "",
        output_text,
        flags=re.IGNORECASE,
    )

    # Extract the first valid numeric score in [0, 1]
    matches = re.findall(
        r"(?<!\d)(?:0\.\d+|1\.0*|0|1)(?!\d)",
        clean_text,
    )

    if not matches:
        raise ValueError(
            "No valid similarity score found in model output: "
            f"{output_text!r}"
        )

    return float(matches[0])


# ============================================================
# 6. EXCEL COLUMN FUNCTIONS
# ============================================================

def get_column_index(
    worksheet,
    header_name: str,
) -> int:

    expected = header_name.strip().lower()

    for index, cell in enumerate(
        worksheet[1],
        start=1,
    ):
        value = cell.value

        if (
            isinstance(value, str)
            and value.strip().lower() == expected
        ):
            return index

    raise ValueError(
        f"Column '{header_name}' was not found "
        "in the workbook header."
    )


def get_or_create_column_index(
    worksheet,
    header_name: str,
) -> int:

    expected = header_name.strip().lower()
    blank_column = None

    for index, cell in enumerate(
        worksheet[1],
        start=1,
    ):
        value = cell.value

        if (
            isinstance(value, str)
            and value.strip().lower() == expected
        ):
            return index

        if (
            blank_column is None
            and (
                value is None
                or (
                    isinstance(value, str)
                    and not value.strip()
                )
            )
        ):
            blank_column = index

    if blank_column is None:
        blank_column = worksheet.max_column + 1

    worksheet.cell(
        row=1,
        column=blank_column,
    ).value = header_name

    return blank_column


def should_generate_score(
    existing_score,
) -> bool:

    if existing_score is None:
        return True

    if (
        isinstance(existing_score, str)
        and not existing_score.strip()
    ):
        return True

    return False


# ============================================================
# 7. LOAD MODEL AND TOKENIZER
# ============================================================

def load_model_and_tokenizer(
    model_id: str,
):

    device = (
        "cuda"
        if torch.cuda.is_available()
        else "cpu"
    )

    print("Loading tokenizer...")

    tokenizer = AutoTokenizer.from_pretrained(
        model_id,
        trust_remote_code=True,
    )

    if (
        tokenizer.pad_token_id is None
        and tokenizer.eos_token_id is not None
    ):
        tokenizer.pad_token = tokenizer.eos_token

    print(
        "Loading model with 4-bit NF4 quantization..."
    )

    quantization_config = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_compute_dtype=torch.float16,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_use_double_quant=True,
    )

    model = AutoModelForCausalLM.from_pretrained(
        model_id,
        device_map={"": 0},
        quantization_config=quantization_config,
        trust_remote_code=True,
    )

    model.eval()

    return model, tokenizer, device


# ============================================================
# 8. SCORE GENERATION
# ============================================================

def generate_score(
    model,
    tokenizer,
    device: str,
    text1: str,
    text2: str,
) -> float:

    prompt = build_prompt(
        text1,
        text2,
    )

    input_text = (
        f"### Instruction:\n"
        f"{prompt}\n\n"
        f"### Response:\n"
        f"সদাদৃশ্য স্কোর: "
    )

    inputs = tokenizer(
        input_text,
        return_tensors="pt",
    )

    inputs = {
        key: value.to(device)
        for key, value in inputs.items()
    }

    generation_kwargs = {
        "max_new_tokens": 15,
        "do_sample": False,
        "pad_token_id": (
            tokenizer.pad_token_id
            if tokenizer.pad_token_id is not None
            else tokenizer.eos_token_id
        ),
    }

    terminators = [
        tokenizer.eos_token_id
    ]

    try:
        eot_token_id = (
            tokenizer.convert_tokens_to_ids(
                "<|eot_id|>"
            )
        )

        if eot_token_id is not None:
            terminators.append(
                eot_token_id
            )

    except Exception:
        pass

    generation_kwargs[
        "eos_token_id"
    ] = terminators

    with torch.inference_mode():

        outputs = model.generate(
            **inputs,
            **generation_kwargs,
        )

    input_length = (
        inputs["input_ids"].shape[-1]
    )

    generated_tokens = outputs[0][
        input_length:
    ]

    response_text = tokenizer.decode(
        generated_tokens,
        skip_special_tokens=True,
    ).strip()

    full_response = (
        "সদাদৃশ্য স্কোর: "
        + response_text
    )

    print(
        f"[Model Says]: {full_response}"
    )

    return extract_score(
        full_response
    )


# ============================================================
# 9. PROCESS WORKBOOK
# ============================================================

def process_workbook(
    workbook_path: str,
    model,
    tokenizer,
    device: str,
    score_column_name: str,
) -> None:

    if not os.path.exists(
        workbook_path
    ):
        raise FileNotFoundError(
            f"File not found: "
            f"{workbook_path}"
        )

    print(
        f"Processing file: "
        f"{workbook_path}"
    )

    workbook = load_workbook(
        workbook_path
    )

    worksheet = workbook.active

    text1_column = get_column_index(
        worksheet,
        "text1",
    )

    text2_column = get_column_index(
        worksheet,
        "text2",
    )

    score_column = (
        get_or_create_column_index(
            worksheet,
            score_column_name,
        )
    )

    total_rows = (
        worksheet.max_row - 1
    )

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):

        existing_score = worksheet.cell(
            row=row_number,
            column=score_column,
        ).value

        if not should_generate_score(
            existing_score
        ):

            print(
                f"Skipping row "
                f"{row_number - 1}/"
                f"{total_rows}: "
                f"existing score "
                f"{existing_score}"
            )

            continue

        text1 = worksheet.cell(
            row=row_number,
            column=text1_column,
        ).value

        text2 = worksheet.cell(
            row=row_number,
            column=text2_column,
        ).value

        if (
            text1 is None
            or text2 is None
        ):
            continue

        try:

            score = generate_score(
                model,
                tokenizer,
                device,
                str(text1),
                str(text2),
            )

            worksheet.cell(
                row=row_number,
                column=score_column,
            ).value = score

            workbook.save(
                workbook_path
            )

            print(
                f"Processed row "
                f"{row_number - 1}/"
                f"{total_rows}: "
                f"{score:.4f}"
            )

        except Exception as error:

            print(
                f"Error on row "
                f"{row_number - 1}: "
                f"{error}"
            )


# ============================================================
# 10. MAIN
# ============================================================

def main():

    model, tokenizer, device = (
        load_model_and_tokenizer(
            MODEL_ID
        )
    )

    try:

        process_workbook(
            WORKBOOK_PATH,
            model,
            tokenizer,
            device,
            SCORE_COLUMN_NAME,
        )

        print(
            "Excel file updated successfully."
        )

    finally:

        del model

        gc.collect()

        if torch.cuda.is_available():
            torch.cuda.empty_cache()


if __name__ == "__main__":
    main()