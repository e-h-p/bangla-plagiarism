import os
import gc
import re
import torch
from openpyxl import load_workbook
from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig

HF_TOKEN = os.environ.get("HF_TOKEN")
if not HF_TOKEN:
    raise RuntimeError(
        "HF_TOKEN is not set. Add it to the environment; "
        "never commit the token to Git."
    )

# ===============================
# CONFIGURATION
# ===============================

WORKBOOK_PATH = "./data/input.xlsx"  # Set this to your dataset workbook path

# ===============================
# MODEL CONFIGURATION
# ===============================

# Supported Model IDs:
# 1) KillerShoaib/llama-3-8b-bangla-4bit
# 2) google/gemma-2-9b-it
# 3) BanglaLLM/BanglaLLama-3-8b-bangla-alpaca-orca-instruct-v0.0.1

MODEL_ID = "google/gemma-2-9b-it"
MODEL_SHORT_NAME = "Gemma_2_9B_IT"

# Save workbook every N rows
SAVE_EVERY_N_ROWS = 10

# Test configuration (set END_ROW to None for all rows)
START_ROW = 2          # Excel row 2 = first data row
END_ROW = None         # None means process all rows
PROMPT_VARIANTS = {

    "P1_FewShot_Anchored": """
Act as an expert Bangla plagiarism evaluator.

Use the following examples as scoring anchors.

Example 1 — No Plagiarism:

Text 1:
বাংলাদেশের কৃষিভিত্তিক অর্থনীতিতে ধান একটি অত্যন্ত গুরুত্বপূর্ণ খাদ্যশস্য। দেশের অধিকাংশ কৃষক বর্ষা মৌসুমে ধান চাষের উপর নির্ভরশীল। ধান উৎপাদন দেশের খাদ্য নিরাপত্তা নিশ্চিত করতে গুরুত্বপূর্ণ ভূমিকা পালন করে।

Text 2:
বাংলাদেশের নদীগুলো বর্ষাকালে অতিরিক্ত পানি বহন করে এবং অনেক অঞ্চলে বন্যার সৃষ্টি হয়। নদীপথ দেশের যোগাযোগ ও বাণিজ্যের জন্যও গুরুত্বপূর্ণ ভূমিকা পালন করে।

Expected score: 0.10


Example 2 — Light Paraphrasing:

Text 1:
পরীক্ষার আগে অধিকাংশ শিক্ষার্থী নিয়মিত পড়াশোনা করে এবং গুরুত্বপূর্ণ বিষয়গুলো বারবার অনুশীলন করে। ভালো ফলাফল অর্জনের জন্য ধারাবাহিক অধ্যবসায় প্রয়োজন।

Text 2:
ভালো ফল করার উদ্দেশ্যে পরীক্ষার পূর্বে অনেক ছাত্রছাত্রী নিয়মিত অধ্যয়ন করে এবং গুরুত্বপূর্ণ অধ্যায়গুলো পুনরায় চর্চা করে। সফলতার জন্য নিয়মিত পরিশ্রম গুরুত্বপূর্ণ।

Expected score: 0.55


Example 3 — Heavy Paraphrasing:

Text 1:
জলবায়ু পরিবর্তনের কারণে বাংলাদেশের উপকূলীয় অঞ্চলে লবণাক্ততার পরিমাণ বৃদ্ধি পাচ্ছে। এর ফলে কৃষিজমির উর্বরতা কমে যাচ্ছে এবং বিশুদ্ধ পানির সংকট দেখা দিচ্ছে।

Text 2:
দীর্ঘমেয়াদি আবহাওয়াগত পরিবর্তনের প্রভাবে বাংলাদেশের সমুদ্র উপকূলবর্তী এলাকাগুলোতে মাটির লবণভাব বৃদ্ধি পাচ্ছে। এতে কৃষি উৎপাদন ব্যাহত হচ্ছে এবং নিরাপদ পানির প্রাপ্যতা হ্রাস পাচ্ছে।

Expected score: 0.78


Example 4 — Direct Copy:

Text 1:
বাংলা ভাষা আমাদের সাংস্কৃতিক পরিচয়ের অন্যতম গুরুত্বপূর্ণ অংশ। এই ভাষার মাধ্যমে আমরা আমাদের ইতিহাস, সাহিত্য এবং ঐতিহ্যকে প্রকাশ করি। আন্তর্জাতিক মাতৃভাষা দিবস বাংলা ভাষার মর্যাদা বিশ্বব্যাপী প্রতিষ্ঠা করেছে।

Text 2:
বাংলা ভাষা আমাদের সাংস্কৃতিক পরিচয়ের অন্যতম গুরুত্বপূর্ণ অংশ। এই ভাষার মাধ্যমে আমরা আমাদের ইতিহাস, সাহিত্য এবং ঐতিহ্যকে প্রকাশ করি। আন্তর্জাতিক মাতৃভাষা দিবস বাংলা ভাষার মর্যাদা বিশ্বব্যাপী প্রতিষ্ঠা করেছে।

Expected score: 0.96


Now evaluate the following pair using the same scoring pattern.

Rules:
- Same topic alone is not plagiarism.
- Preserved central meaning, claim, and facts indicate plagiarism.
- Synonym, morphology, lexical, or word-order changes still count as plagiarism if meaning is preserved.
- Topic similarity with different claim or facts should receive a low score.

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

    "P2_Contradiction_Aware": """
Act as a strict Bangla plagiarism evaluator.

Your task is to determine whether Text 2 preserves the same meaning and factual direction as Text 1.

Before assigning a high plagiarism score, carefully verify:
- the central conclusion is preserved,
- the factual meaning is preserved,
- the logical direction is consistent,
- the texts do not contradict each other.

Important:
- Same topic alone is NOT plagiarism.
- If the topic is similar but the conclusion or factual claim is opposite, assign a low score.
- Contradictory or independently written texts should receive low scores even if they share keywords or topic overlap.

Scoring guideline:
0.00–0.20 = unrelated or contradictory
0.21–0.40 = same topic only
0.41–0.65 = partial semantic overlap
0.66–0.85 = strong paraphrasing
0.86–1.00 = near-copy or direct plagiarism

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

    "P3_Hard_Negative_Aware": """
Act as a strict Bangla plagiarism evaluator.

Your main goal is to avoid false positives.

Many text pairs may appear similar because they discuss the same topic. However, topical similarity alone does NOT indicate plagiarism unless the same meaning, factual content, and central claim are preserved.

Assign a high score only if Text 2 appears derived from Text 1 through:
- paraphrasing,
- synonym replacement,
- word-order modification,
- lexical rewriting,
- morphological transformation,
- or mixed semantic rewriting.

Assign a low score if:
- the topic is similar but the claim is different,
- the facts are different,
- the conclusion is different,
- the second text appears independently written.

Scoring:
0.00–0.20 = unrelated
0.21–0.40 = same topic only
0.41–0.65 = light paraphrasing
0.66–0.85 = heavy paraphrasing
0.86–1.00 = direct plagiarism

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

"P4_Bangla_Instruction": """
একজন বিশেষজ্ঞ বাংলা চৌর্যবৃত্তি মূল্যায়নকারী হিসেবে কাজ করুন।

শুধু একই শব্দ ব্যবহৃত হয়েছে কি না, তা না দেখে লেখা দুটির অর্থ, মূল বক্তব্য এবং তথ্যগত মিলের ভিত্তিতে চৌর্যবৃত্তির মাত্রা নির্ধারণ করুন।

গুরুত্বপূর্ণ নির্দেশনা:
- শুধু একই বিষয় নিয়ে লেখা হলেই সেটি চৌর্যবৃত্তি নয়।
- Text 2-তে Text 1-এর মূল অর্থ, দাবি এবং গুরুত্বপূর্ণ তথ্য বজায় থাকলে সেটি চৌর্যবৃত্তির ইঙ্গিত দেয়।
- সমার্থক শব্দের ব্যবহার, শব্দের রূপ পরিবর্তন, শব্দের ক্রম পরিবর্তন, বাক্য পুনর্গঠন বা ভিন্নভাবে প্রকাশ করা হলেও অর্থ একই থাকলে সেটিকে পুনর্লিখন হিসেবে বিবেচনা করুন।
- বিষয় একই হলেও দাবি, তথ্য, কারণ, ফলাফল বা সিদ্ধান্ত ভিন্ন হলে কম স্কোর দিন।
- লেখা দুটি পরস্পরবিরোধী বা স্বাধীনভাবে লেখা হলে কম স্কোর দিন।
- শব্দের মিল কম হলেও মূল অর্থ ও তথ্য একই থাকলে বেশি স্কোর দেওয়া যেতে পারে।

স্কোর নির্ধারণের মানদণ্ড:
0.00–0.20 = সম্পর্কহীন, পরস্পরবিরোধী বা স্বাধীন লেখা
0.21–0.40 = শুধু বিষয়গত মিল, কিন্তু দাবি বা তথ্য ভিন্ন
0.41–0.65 = আংশিক অর্থগত মিল বা হালকা পুনর্লিখন
0.66–0.85 = মূল অর্থ ও তথ্য বজায় রেখে ব্যাপক পুনর্লিখন
0.86–1.00 = প্রায় হুবহু অনুলিপি, সম্পূর্ণ অর্থগত সমতা বা সরাসরি চৌর্যবৃত্তি

Text 1: {text1}

Text 2: {text2}

শুধু 0.00 থেকে 1.00-এর মধ্যে একটি সংখ্যাসূচক স্কোর দিন। কোনো ব্যাখ্যা, শব্দ বা অতিরিক্ত লেখা দেবেন না। স্কোরটি অবশ্যই ইংরেজি অঙ্ক এবং দশমিক বিন্দু ব্যবহার করে লিখবেন।
""",

    "P5_NLI_Semantic": """
Act as a semantic plagiarism evaluator using Natural Language Inference principles.

Determine whether Text 2 semantically entails, preserves, or reproduces the meaning of Text 1 through paraphrasing or rewriting.

Assign a high score only if:
- the same meaning is preserved,
- the same factual information is preserved,
- the same central claim is preserved,
- the semantic intent remains equivalent.

Assign a low score if:
- the texts only share the same topic,
- the factual direction changes,
- the conclusions differ,
- or contradictory information appears.

Important:
Surface-level lexical overlap alone should not determine the score.

Scoring:
0.00–0.20 = unrelated
0.21–0.40 = same topic only
0.41–0.65 = partial semantic overlap
0.66–0.85 = strong semantic paraphrasing
0.86–1.00 = semantic equivalence or direct plagiarism

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

    "P6_MultiDimensional": """
Act as an expert Bangla plagiarism evaluator.

Internally evaluate the pair across multiple dimensions:

1. Semantic overlap:
Does Text 2 preserve the overall meaning of Text 1?

2. Central claim similarity:
Do both texts express the same main claim or message?

3. Factual consistency:
Are the same key facts, entities, quantities, causes, or conclusions preserved?

4. Paraphrase intensity:
Does Text 2 appear rewritten using synonyms, word-order changes, morphology changes, lexical substitutions, or sentence restructuring?

5. Contradiction possibility:
Do the texts contradict each other or express different factual directions?

6. Topic-only similarity:
Are the texts only about the same topic without preserving the same claim?

After checking all dimensions internally, combine them into one final plagiarism score.

Important:
- Same topic alone is NOT plagiarism.
- High lexical overlap is not enough if the meaning differs.
- Low lexical overlap can still be plagiarism if the same meaning is preserved.
- Contradictory or independently written texts should receive low scores.
- Heavy paraphrasing should receive high scores if meaning and facts are preserved.

Scoring:
0.00–0.20 = unrelated
0.21–0.40 = same topic only
0.41–0.65 = partial or light paraphrasing
0.66–0.85 = strong or heavy paraphrasing
0.86–1.00 = near-copy, semantic equivalence, or direct plagiarism

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

    "P7_Calibrated_Conservative": """
Act as a conservative Bangla plagiarism evaluator.

Use the full score range carefully and avoid overestimating plagiarism.

Do NOT assign a high score unless Text 2 clearly preserves the same central meaning, claim, and factual information from Text 1.

Scoring rubric:
0.00–0.20 = unrelated or independent writing
0.21–0.40 = same broad topic only, but different claim or facts
0.41–0.65 = partial semantic overlap or light paraphrasing
0.66–0.85 = strong paraphrasing with preserved meaning and facts
0.86–1.00 = direct copy, near-copy, or semantic equivalence

Conservative rules:
- Same topic alone should stay below 0.40.
- Different facts or conclusions should stay below 0.40.
- Partial overlap without full meaning preservation should usually stay below 0.65.
- Scores above 0.85 should be used only for near-copy or almost identical meaning.
- If uncertain between two score ranges, choose the lower range.

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
""",

    "P8_Direct_Short": """
Evaluate whether Text 2 plagiarizes Text 1.

Focus on semantic meaning, not only word overlap.

Score:
0.00 = no plagiarism
1.00 = direct plagiarism

Text 1: {text1}

Text 2: {text2}

Return only one numeric score between 0.00 and 1.00.
"""
}


# ===============================
# HELPER FUNCTIONS
# ===============================

def build_prompt(prompt_template: str, text1: str, text2: str) -> str:
    return prompt_template.format(text1=text1, text2=text2)


def extract_score(output_text: str) -> float:
    # Convert Bangla digits to English digits
    digit_map = str.maketrans("০১২৩৪৫৬৭৮৯", "0123456789")
    normalized_text = output_text.translate(digit_map)

    clean_text = re.sub(
        r"(text|পাঠ্য|ট্যাক্সট|টেক্সট)\s*[12]",
        "",
        normalized_text,
        flags=re.IGNORECASE
    )

    matches = re.findall(
        r"(?<!\d)(?:0\.\d+|1\.0*|0|1)(?!\d)",
        clean_text
    )

    if not matches:
        raise ValueError(
            f"No valid score found in model output: {output_text!r}"
        )

    score = float(matches[0])

    if not 0.0 <= score <= 1.0:
        raise ValueError(f"Extracted score out of range: {score}")

    return score


def get_column_index(worksheet, header_name: str) -> int:
    expected = header_name.strip().lower()

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value

        if isinstance(value, str) and value.strip().lower() == expected:
            return index

    raise ValueError(f"Column '{header_name}' was not found in the workbook header.")


def get_or_create_column_index(worksheet, header_name: str) -> int:
    expected = header_name.strip().lower()
    blank_column = None

    for index, cell in enumerate(worksheet[1], start=1):
        value = cell.value

        if isinstance(value, str) and value.strip().lower() == expected:
            return index

        if blank_column is None and (value is None or (isinstance(value, str) and not value.strip())):
            blank_column = index

    if blank_column is None:
        blank_column = worksheet.max_column + 1

    worksheet.cell(row=1, column=blank_column).value = header_name
    return blank_column


def should_generate_score(existing_score) -> bool:
    if existing_score is None:
        return True

    if isinstance(existing_score, str):
        return existing_score.strip() == ""

    return False


# ===============================
# MODEL LOADING
# ===============================

def load_model_and_tokenizer(model_id: str):
    if not torch.cuda.is_available():
        raise RuntimeError("CUDA GPU is not available. Please use a GPU runtime.")

    device = "cuda"

    print("Loading tokenizer...")
    tokenizer = AutoTokenizer.from_pretrained(
        model_id,
        trust_remote_code=True,
        token=os.environ.get("HF_TOKEN")
    )

    if tokenizer.pad_token_id is None and tokenizer.eos_token_id is not None:
        tokenizer.pad_token = tokenizer.eos_token

    print("Loading model with 4-bit quantization...")
    quantization_config = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_compute_dtype=torch.float16,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_use_double_quant=True
    )

    model = AutoModelForCausalLM.from_pretrained(
        model_id,
        device_map={"": 0},
        quantization_config=quantization_config,
        trust_remote_code=True,
        token=os.environ.get("HF_TOKEN")
    )

    model.eval()
    return model, tokenizer, device


# ===============================
# SCORE GENERATION
# ===============================

def generate_score(
    model,
    tokenizer,
    device: str,
    text1: str,
    text2: str,
    prompt_template: str
) -> float:
    prompt = build_prompt(prompt_template, text1, text2)

    input_text = f"""### Instruction:
{prompt}

### Response:
Similarity Score: """

    inputs = tokenizer(
        input_text,
        return_tensors="pt",
        truncation=True,
        max_length=4096
    )

    inputs = {k: v.to(device) for k, v in inputs.items()}

    generation_kwargs = {
        "max_new_tokens": 12,
        "do_sample": False,
        "pad_token_id": tokenizer.pad_token_id if tokenizer.pad_token_id is not None else tokenizer.eos_token_id,
        "eos_token_id": tokenizer.eos_token_id,
    }

    with torch.inference_mode():
        outputs = model.generate(**inputs, **generation_kwargs)

    input_length = inputs["input_ids"].shape[-1]
    generated_tokens = outputs[0][input_length:]
    response_text = tokenizer.decode(generated_tokens, skip_special_tokens=True).strip()

    full_response = "Similarity Score: " + response_text
    print(f"[Model Says]: {full_response}")

    return extract_score(full_response)


# ===============================
# WORKBOOK PROCESSING
# ===============================

def process_prompt_variant(
    workbook_path: str,
    model,
    tokenizer,
    device: str,
    prompt_name: str,
    prompt_template: str
) -> None:
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    score_column_name = f"{MODEL_SHORT_NAME}_{prompt_name}_Score"

    print("\n" + "=" * 80)
    print(f"Processing prompt: {prompt_name}")
    print(f"Score column: {score_column_name}")
    print("=" * 80)

    workbook = load_workbook(workbook_path)
    worksheet = workbook.active

    text1_column = get_column_index(worksheet, "text1")
    text2_column = get_column_index(worksheet, "text2")
    score_column = get_or_create_column_index(worksheet, score_column_name)

    last_row = worksheet.max_row if END_ROW is None else min(END_ROW, worksheet.max_row)

    processed_since_save = 0

    for row_number in range(START_ROW, last_row + 1):
        existing_score = worksheet.cell(row=row_number, column=score_column).value

        if not should_generate_score(existing_score):
            print(f"Skipping row {row_number - 1}/{last_row - 1}: existing score = {existing_score}")
            continue

        text1 = worksheet.cell(row=row_number, column=text1_column).value
        text2 = worksheet.cell(row=row_number, column=text2_column).value

        if text1 is None or text2 is None:
            print(f"Skipping row {row_number - 1}: empty text1/text2")
            continue

        try:
            score = generate_score(
                model=model,
                tokenizer=tokenizer,
                device=device,
                text1=str(text1),
                text2=str(text2),
                prompt_template=prompt_template
            )

            worksheet.cell(row=row_number, column=score_column).value = score
            processed_since_save += 1

            print(
                f"Processed prompt={prompt_name}, "
                f"row={row_number - 1}/{last_row - 1}, "
                f"score={score:.4f}"
            )

            if processed_since_save >= SAVE_EVERY_N_ROWS:
                workbook.save(workbook_path)
                processed_since_save = 0
                print("Workbook saved.")

        except Exception as e:
            print(f"Error on prompt={prompt_name}, row={row_number - 1}: {e}")

    workbook.save(workbook_path)
    print(f"Finished prompt {prompt_name}. Workbook saved.")


# ===============================
# MAIN
# ===============================

def main():
    model, tokenizer, device = load_model_and_tokenizer(MODEL_ID)

    try:
        for prompt_name, prompt_template in PROMPT_VARIANTS.items():
            process_prompt_variant(
                workbook_path=WORKBOOK_PATH,
                model=model,
                tokenizer=tokenizer,
                device=device,
                prompt_name=prompt_name,
                prompt_template=prompt_template
            )

        print("\nAll 8 prompt variants processed successfully.")
        print("Check your Excel file in Google Drive.")

    finally:
        del model
        gc.collect()
        torch.cuda.empty_cache()


if __name__ == "__main__":
    main()
